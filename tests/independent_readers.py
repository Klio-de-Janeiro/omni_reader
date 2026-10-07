"""Reopen generated editor outputs with libraries independent of the app engines."""
from pathlib import Path
import csv
import json
import wave

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


def main():
    """Validate values, page rotation, audio frames and Unicode after npm test."""
    root = Path(__file__).resolve().parents[1] / 'test-results' / 'edited'
    expected = '  Клио <новый> & "текст" 😀  '
    assert Document(root / 'sample-edited.docx').paragraphs[0].text == expected
    assert Presentation(root / 'sample-edited.pptx').slides[0].shapes.title.text == expected
    book = load_workbook(root / 'sample-edited.xlsx', data_only=False)
    assert book.worksheets[0]['A2'].value == '<&Клио> 😀'
    assert book.worksheets[0]['B2'].value == 123.5
    assert book.worksheets[0]['D2'].value == '=SUM(B2:B3)'
    assert book.worksheets[0]['E2'].data_type == 's'
    assert book.worksheets[1]['D301'].value is True
    pdf = PdfReader(root / 'sample-edited.pdf')
    assert len(pdf.pages) == 1 and pdf.pages[0].rotation == 90
    assert 'OMNI PDF - PAGE 1' in pdf.pages[0].extract_text()
    with wave.open(str(root / 'sample-edited.wav'), 'rb') as audio:
        assert audio.getnframes() == 12000
        assert audio.getframerate() == 16000 and audio.getsampwidth() == 2
    with (root / 'sample-edited.csv').open(encoding='utf-8-sig', newline='') as stream:
        assert list(csv.reader(stream, delimiter=';'))[1][1] == 'Новое; "значение"\nстрока'
    rich = root.parent / 'rich'
    word = Document(rich / 'formatted.docx')
    run = next(run for run in word.paragraphs[0].runs if run.text == 'Клио <&> 😀')
    assert run.bold and run.italic and run.underline and run.font.strike
    presentation = Presentation(rich / 'formatted.pptx')
    run = next(run for run in presentation.slides[0].shapes.title.text_frame.paragraphs[0].runs if run.text == 'Клио <&> 😀')
    assert run.font.bold and run.font.italic and run.font.underline and run._r.rPr.get('strike') == 'sngStrike'
    word = Document(rich / 'page.docx')
    assert next(p for p in word.paragraphs if p.text == 'Новая страница').paragraph_format.page_break_before
    assert len(Presentation(rich / 'sections.pptx').slides) == 4
    run = Document(rich / 'no-styles.docx').paragraphs[0].runs[0]
    assert run.bold is False and run.italic is False and run.underline is False and run.font.strike is False
    run = Presentation(rich / 'no-styles.pptx').slides[0].shapes.title.text_frame.paragraphs[0].runs[0]
    assert run.font.bold is False and run.font.italic is False and run.font.underline is False and run._r.rPr.get('strike') == 'noStrike'
    modes = root.parent / 'modes'
    assert Document(modes / 'layout.docx').paragraphs[0].text == 'Документ ✓'
    presentation = Presentation(modes / 'layout.pptx')
    run = next(run for run in presentation.slides[1].shapes.title.text_frame.paragraphs[0].runs if run.text == 'Презентация ✓')
    assert run.font.bold and run.font.italic and run.font.underline and run._r.rPr.get('strike') == 'sngStrike'
    result = {'version': '0.3.0-laptop-inline-modes', 'status': 'passed', 'formats': ['docx', 'pptx', 'xlsx', 'pdf', 'wav', 'csv'], 'richExports': 'passed', 'modeSwitchExports': 'passed'}
    (root.parent / 'independent-readers.json').write_text(json.dumps(result, indent=2))
    print(json.dumps(result))


if __name__ == '__main__':
    main()
