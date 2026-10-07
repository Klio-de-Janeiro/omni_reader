"""Generate small, deterministic fixtures for viewer acceptance tests."""
from pathlib import Path
import math
import struct
import wave

from docx import Document
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / 'tests' / 'fixtures'
FIXTURES.mkdir(parents=True, exist_ok=True)

image = Image.new('RGB', (1000, 650), '#e8edff')
draw = ImageDraw.Draw(image)
draw.rounded_rectangle((120, 120, 880, 530), 50, fill='#5551eb')
draw.text((400, 310), 'OMNI TEST IMAGE', fill='white', font_size=32)
image.save(FIXTURES / 'sample.png')
image.save(FIXTURES / 'sample.jpg')

pdf = canvas.Canvas(str(FIXTURES / 'sample.pdf'))
for number in (1, 2):
    pdf.setFont('Helvetica', 28)
    pdf.drawString(72, 720, f'OMNI PDF - PAGE {number}')
    pdf.setFont('Helvetica', 16)
    pdf.drawString(72, 670, 'A local document viewer acceptance test.')
    pdf.drawImage(str(FIXTURES / 'sample.png'), 72, 330, 380, 247)
    pdf.showPage()
pdf.save()

document = Document()
document.add_heading('OMNI DOCX TEST', 0)
document.add_paragraph('Документ на русском: проверка текста, таблицы и картинки.')
table = document.add_table(rows=2, cols=2)
table.cell(0, 0).text = 'Формат'
table.cell(0, 1).text = 'Статус'
table.cell(1, 0).text = 'DOCX'
table.cell(1, 1).text = 'Работает локально'
document.add_picture(str(FIXTURES / 'sample.png'), width=Inches(4))
document.add_page_break()
document.add_heading('Вторая страница', 1)
document.add_paragraph('Проверка перехода между страницами.')
document.save(FIXTURES / 'sample.docx')

presentation = Presentation()
for number in (1, 2):
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = f'OMNI PPTX — слайд {number}'
    slide.shapes.add_picture(str(FIXTURES / 'sample.png'), Inches(1), Inches(2), width=Inches(5))
presentation.save(FIXTURES / 'sample.pptx')

with wave.open(str(FIXTURES / 'sample.wav'), 'wb') as audio:
    audio.setnchannels(1)
    audio.setsampwidth(2)
    audio.setframerate(16000)
    samples = [int(2000 * math.sin(2 * math.pi * 440 * n / 16000)) for n in range(32000)]
    audio.writeframes(struct.pack('<' + 'h' * len(samples), *samples))

print('Six document/image/audio fixtures generated. Existing app icons and table fixtures are unchanged.')
